"""Turn any uploaded source into a DOCX the whitepaper builder can read.

    py -3.13 source_to_docx.py <source> <output.docx>

The wp pipeline (whitepaper/build_whitepaper.py) only reads DOCX: it takes the page layout from
the heading styles and the cover title from the first Title/Heading1 paragraph. This script lets a
task start from the other formats the console accepts:

    .docx            copied as is
    .pdf             text + pictures, in reading order (headings inferred from font size)
    .pptx            one section per slide: title -> Heading 1, text -> paragraphs, pictures kept
    .md / .txt       '#' lines or a short first line -> headings, blank line -> new paragraph
    .png .jpg .jpeg  the picture alone, under a title taken from the file name
    .svg             rasterised with PyMuPDF, then as a picture

Text is copied verbatim, in the source language. Pictures are never turned into text; a source
that is ONLY a picture has no words to copy, so the caller (Claude) must read it and write the
text into the DOCX by hand.
"""
import os
import re
import shutil
import sys
import tempfile

from docx import Document
from docx.shared import Inches

ZW = re.compile('[​‌‍﻿]')
LIST_START = re.compile(r'^\s*(\d+[.)]|[•\-–▪●◦])\s+')
MAX_W = 6.0  # inches, the text column of a portrait A4/Letter page


def clean(t):
    return ZW.sub('', t).replace('\xa0', ' ')


def new_doc(title):
    doc = Document()
    doc.add_paragraph(title, style='Title')
    return doc


def add_line(doc, text):
    """Numbered / bulleted lines keep their list look; everything else is a plain paragraph."""
    m = LIST_START.match(text)
    if not m:
        doc.add_paragraph(text)
    elif m.group(1)[0].isdigit():
        doc.add_paragraph(text[m.end():], style='List Number')
    else:
        doc.add_paragraph(text[m.end():], style='List Bullet')


def add_picture(doc, path, width_in=MAX_W):
    doc.add_picture(path, width=Inches(min(width_in, MAX_W)))


# ── PDF ─────────────────────────────────────────────────────────────────────
def from_pdf(src, out):
    import fitz
    d = fitz.open(src)
    tmp = tempfile.mkdtemp()
    sizes = {}
    for p in d:
        for b in p.get_text('dict')['blocks']:
            for l in b.get('lines', []):
                for s in l['spans']:
                    if clean(s['text']).strip():
                        sizes[round(s['size'], 1)] = sizes.get(round(s['size'], 1), 0) + len(s['text'])
    body = max(sizes, key=sizes.get) if sizes else 11
    big = sorted({s for s in sizes if s >= body * 1.15}, reverse=True)  # heading sizes, largest first

    # items: (page, y, kind, payload) — text paragraphs and picture clips, in reading order
    items = []
    for p in d:
        lines = []
        for b in p.get_text('dict')['blocks']:
            for l in b.get('lines', []):
                txt = clean(''.join(s['text'] for s in l['spans'])).strip()
                if txt:
                    lines.append([l['bbox'][1], l['bbox'][3] - l['bbox'][1], round(l['spans'][0]['size'], 1), txt, l['bbox'][0]])
        lines.sort(key=lambda r: (round(r[0] / 3), r[4]))  # row first, then left to right
        rows = []
        for r in lines:
            if rows and abs(r[0] - rows[-1][0]) < r[1] * 0.5:
                rows[-1][3] += ' ' + r[3]
            else:
                rows.append(r)
        lines = rows
        cur = None
        for y, h, size, txt, _x in lines:
            new = (cur is None or size != cur['size'] or LIST_START.match(txt)
                   or y - cur['y1'] > h * 0.9)
            if new:
                cur = {'y': y, 'y1': y + h, 'size': size, 'text': txt}
                items.append([p.number, y, 'text', cur])
            else:
                cur['text'] += ' ' + txt
                cur['y1'] = y + h
        for k, info in enumerate(sorted(p.get_image_info(), key=lambda i: i['bbox'][1])):
            fn = os.path.join(tmp, f'p{p.number + 1}_{k}.png')
            # render the visible area (what the page shows), not the raw image, which may be cropped
            p.get_pixmap(clip=fitz.Rect(info['bbox']), dpi=200).save(fn)
            w = (info['bbox'][2] - info['bbox'][0]) / 72
            items.append([p.number, info['bbox'][1], 'img', (fn, w)])
    items.sort(key=lambda i: (i[0], i[1]))

    texts = [i for i in items if i[2] == 'text']
    title_item = next((i for i in texts if big and i[3]['size'] == big[0]), None) or (texts[0] if texts else None)
    doc = new_doc(title_item[3]['text'] if title_item else os.path.splitext(os.path.basename(src))[0])
    skip = {id(title_item)} if title_item else set()
    for it in items:
        if id(it) in skip:
            continue
        if it[2] == 'img':
            add_picture(doc, *it[3])
        else:
            c = it[3]
            if c['size'] == (title_item[3]['size'] if title_item else -1) and it is not title_item:
                doc.add_heading(c['text'], level=1)
            elif c['size'] in big:
                rest = [b for b in big if not title_item or b != title_item[3]['size']]
                doc.add_heading(c['text'], level=min(rest.index(c['size']) + 1, 3))
            else:
                add_line(doc, c['text'])
    doc.save(out)


# ── PPTX ────────────────────────────────────────────────────────────────────
def from_pptx(src, out):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(src)
    tmp = tempfile.mkdtemp()
    first = prs.slides[0] if len(prs.slides) else None
    t0 = first.shapes.title.text_frame.text.strip() if first is not None and first.shapes.title is not None else ''
    doc = new_doc(t0 or os.path.splitext(os.path.basename(src))[0])
    for n, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title
        if title is not None and title.text_frame.text.strip() and not (n == 1 and t0):
            doc.add_heading(clean(title.text_frame.text.strip()), level=1)
        for sh in slide.shapes:
            if sh == title:
                continue
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = clean(''.join(r.text for r in para.runs)).strip()
                    if t:
                        add_line(doc, t) if LIST_START.match(t) else doc.add_paragraph(t)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                fn = os.path.join(tmp, f's{n}_{sh.shape_id}.{sh.image.ext}')
                open(fn, 'wb').write(sh.image.blob)
                add_picture(doc, fn, sh.width / 914400)
            elif sh.has_table if hasattr(sh, 'has_table') else False:
                rows = sh.table.rows
                tbl = doc.add_table(rows=len(rows), cols=len(sh.table.columns))
                tbl.style = 'Table Grid'
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row.cells):
                        tbl.cell(r, c).text = clean(cell.text)
    doc.save(out)


# ── MD / TXT ────────────────────────────────────────────────────────────────
def from_text(src, out):
    raw = open(src, encoding='utf-8-sig', errors='replace').read()
    blocks = [b.strip('\n') for b in re.split(r'\n\s*\n', raw) if b.strip()]
    title = None
    if blocks:
        first = clean(blocks[0]).strip()
        title = re.sub(r'^#+\s*', '', first) if '\n' not in first and len(first) < 120 else None
    doc = new_doc(title or os.path.splitext(os.path.basename(src))[0])
    for b in blocks[1 if title else 0:]:
        mixed = any(LIST_START.match(x) or x.lstrip().startswith('#') for x in b.split('\n'))
        for line in b.split('\n') if mixed else [' '.join(b.split('\n'))]:
            line = clean(line).strip()
            m = re.match(r'^(#+)\s+(.*)', line)
            if m:
                doc.add_heading(m.group(2), level=min(len(m.group(1)), 3) if len(m.group(1)) > 1 else 1)
            elif line:
                add_line(doc, line)
    doc.save(out)


# ── Images ──────────────────────────────────────────────────────────────────
def from_image(src, out):
    path = src
    if src.lower().endswith('.svg'):
        import fitz
        svg = fitz.open(src)
        path = os.path.join(tempfile.mkdtemp(), 'svg.png')
        svg[0].get_pixmap(dpi=200).save(path)
    doc = new_doc(os.path.splitext(os.path.basename(src))[0])
    add_picture(doc, path)
    doc.save(out)
    print('NOTE: image-only source — no text was extracted. Read the picture and write its text into '
          'the DOCX (same language) before building.', file=sys.stderr)


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
    ext = os.path.splitext(src)[1].lower()
    if ext == '.docx':
        shutil.copyfile(src, out)
    elif ext == '.pdf':
        from_pdf(src, out)
    elif ext == '.pptx':
        from_pptx(src, out)
    elif ext in ('.md', '.markdown', '.txt'):
        from_text(src, out)
    elif ext in ('.png', '.jpg', '.jpeg', '.svg'):
        from_image(src, out)
    else:
        sys.exit(f'Unsupported source type: {ext}')
    print('wrote', out)


if __name__ == '__main__':
    main()
